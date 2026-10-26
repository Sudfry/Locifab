# -*- coding: utf-8 -*-
"""生成 LocalCraft 参赛项目文档（Document/main.docx）。

结构与要求对齐《关于举办杭州电子科技大学第二届 AIGC 应用创新大赛的通知》：
项目背景、解决方案、技术实现路径、团队分工、创新点与价值、总结与展望，
附录给出 AI 工具使用、关键 Prompt、生成或开发过程、迭代优化、测试验证与人工核查情况。
"""

import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "Document", "main.docx")
SHOT = os.path.join(ROOT, "docs", "screenshots")

SONG = "宋体"
HEI = "黑体"
KAI = "楷体"
LATIN = "Times New Roman"

ACCENT = RGBColor(0x2F, 0x4F, 0x43)  # 深墨绿，与作品界面主色呼应


# ---------------------------------------------------------------- 基础工具


def set_style_font(style, east, size, bold=False, latin=LATIN, color=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def run_font(run, east=SONG, size=12, bold=False, latin=LATIN, color=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def first_line_indent(paragraph, chars=200, twips=480):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), str(twips))


def left_indent(paragraph, cm, hanging=0.0):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), str(int(cm * 567)))
    if hanging:
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), str(int(-hanging * 567)))


def body(text, indent=True, size=12, bold_prefix=None, space_after=0, align=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    pf = p.paragraph_format
    pf.line_spacing_ratio = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if bold_prefix:
        r = p.add_run(bold_prefix)
        run_font(r, SONG, size, bold=True)
    r = p.add_run(text)
    run_font(r, SONG, size)
    if indent:
        first_line_indent(p)
    return p


def bullet(text, size=12, bold_prefix=None, level=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_ratio = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    left_indent(p, 0.74 + 0.74 * level, hanging=0.0)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        run_font(r, SONG, size, bold=True)
    r = p.add_run(text)
    run_font(r, SONG, size)
    return p


def dash_item(text, size=12, bold_prefix=None, indent_cm=0.74):
    """以破折号开头的列表项，避免使用自动编号（跨 Word/WPS 更稳定）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_ratio = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    left_indent(p, indent_cm + 0.6, hanging=0.6)
    ppr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:pos"), str(int((indent_cm + 0.6) * 567)))
    tabs.append(tab)
    ppr.append(tabs)
    r = p.add_run("—\t")
    run_font(r, SONG, size)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        run_font(r, SONG, size, bold=True)
    r = p.add_run(text)
    run_font(r, SONG, size)
    return p


def heading(text, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading %d" % level]
    pf = p.paragraph_format
    if level == 1:
        pf.space_before = Pt(16)
        pf.space_after = Pt(8)
        pf.line_spacing_ratio = 1.5
    elif level == 2:
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.line_spacing_ratio = 1.4
    else:
        pf.space_before = Pt(8)
        pf.space_after = Pt(2)
        pf.line_spacing_ratio = 1.4
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    if level == 1:
        run_font(r, HEI, 16, bold=True)
    elif level == 2:
        run_font(r, HEI, 14, bold=True)
    else:
        run_font(r, HEI, 12, bold=True)
    return p


def caption(text, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    pf.line_spacing_ratio = 1.2
    r = p.add_run(text)
    run_font(r, SONG, 10.5, bold=True)
    if note:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf2 = p2.paragraph_format
        pf2.space_before = Pt(0)
        pf2.space_after = Pt(8)
        pf2.line_spacing_ratio = 1.0
        r2 = p2.add_run(note)
        run_font(r2, SONG, 9)
    return p


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def cell_text(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(1.5)
    pf.space_after = Pt(1.5)
    pf.line_spacing_ratio = 1.15
    r = p.add_run(text)
    run_font(r, SONG, size, bold=bold)


def add_table(headers, rows, widths, size=10.5, header_fill="E8EDEA", first_col_center=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    tblheader = OxmlElement("w:tblHeader")
    trpr.append(tblheader)
    for i, h in enumerate(headers):
        cell_text(hdr.cells[i], h, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr.cells[i], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 and first_col_center) else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[i], val, size=size, align=align)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # 表格行不跨页断裂
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
    return table


def figure(filename, width_cm=14.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(2)
    pf.line_spacing_ratio = 1.0
    p.add_run().add_picture(os.path.join(SHOT, filename), width=Cm(width_cm))


def prompt_block(lines, size=10.5, indent=0.6):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_ratio = 1.25
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        left_indent(p, indent)
        r = p.add_run(line if line else " ")
        run_font(r, SONG, size)
        ppr = p._p.get_or_add_pPr()
        if line:
            pass
    # 块后留白
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing_ratio = 0.6


def add_page_field(paragraph):
    run = paragraph.add_run()
    run_font(run, SONG, 9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, txt, fld_end):
        run._r.append(el)


def page_break():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


# ---------------------------------------------------------------- 文档骨架

doc = Document()

# 页面：A4，上下 2.54cm，左右 3.17cm
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.17)
sec.right_margin = Cm(3.17)
sec.header_distance = Cm(1.5)
sec.footer_distance = Cm(1.5)

# 样式
normal = doc.styles["Normal"]
set_style_font(normal, SONG, 12)
normal.paragraph_format.line_spacing_ratio = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
for lvl, size in ((1, 16), (2, 14), (3, 12)):
    st = doc.styles["Heading %d" % lvl]
    set_style_font(st, HEI, size, bold=True, color=RGBColor(0x00, 0x00, 0x00))
    st.paragraph_format.keep_with_next = True

# 核心属性
cp = doc.core_properties
cp.title = "LocalCraft 本地文件工坊 · 项目文档"
cp.subject = "杭州电子科技大学第二届 AIGC 应用创新大赛 · 赛题一"
cp.author = "LocalCraft 团队"
cp.keywords = "本地文件处理；可回退事务；SimHash；离线软件；AIGC 辅助开发"
cp.comments = "项目文档（Word 格式），符合大赛作品要求"

# ================================================================ 封面
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(46)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("杭州电子科技大学第二届 AIGC 应用创新大赛")
run_font(r, HEI, 15, bold=False)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("赛题一 · AI 辅助开发 · 创意产品原型")
run_font(r, KAI, 12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("项  目  文  档")
run_font(r, HEI, 34, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(28)
r = p.add_run("Project  Documentation")
run_font(r, LATIN, 12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("LocalCraft · 本地文件工坊")
run_font(r, HEI, 22, bold=True, color=ACCENT)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
r = p.add_run("一款完全离线的 Windows 本地文件处理工具")
run_font(r, SONG, 13)

# 封面信息表（无边框）
info_rows = [
    ("团队名称", "　　　　　　　　　　（请填写队名）"),
    ("团队成员", "队长　　　　　（学号　　　　　　）"),
    ("", "成员　　　　　（学号　　　　　　）"),
    ("", "成员　　　　　（学号　　　　　　）"),
    ("", "成员　　　　　（学号　　　　　　）"),
    ("", "成员　　　　　（学号　　　　　　）"),
    ("指导教师", "　　　　　　　　　　（请填写指导教师）"),
    ("所在班级", "　　　　　　　　　　（教学班）"),
    ("完成日期", "2026 年 11 月"),
]
tbl = doc.add_table(rows=0, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in info_rows:
    cells = tbl.add_row().cells
    cell_text(cells[0], label, size=12, align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(cells[1], value, size=12, align=WD_ALIGN_PARAGRAPH.LEFT)
    cells[0].width = Cm(3.6)
    cells[1].width = Cm(8.6)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(40)
r = p.add_run("本作品全部功能在本地完成，不联网、不上传任何文件内容")
run_font(r, KAI, 10.5)

# ================================================================ 摘要
new = doc.add_section(WD_SECTION.NEW_PAGE)
new.page_width = Cm(21.0)
new.page_height = Cm(29.7)
new.top_margin = Cm(2.54)
new.bottom_margin = Cm(2.54)
new.left_margin = Cm(3.17)
new.right_margin = Cm(3.17)

heading("摘　要", 1)

body(
    "LocalCraft（本地文件工坊）是一款面向学生与年轻办公人群的 Windows 桌面文件处理工具，"
    "聚焦“每天都在发生、却长期没人认真解决”的本地文件整理问题。团队通过 27 份有效问卷确认了四个真实缺口："
    "批量操作缺少可回退机制、跨厂商 RAW 格式支持不足、内容级去重缺位、中文场景导出质量差。"
)
body(
    "作品的解决思路是把“安全”做成底层能力：所有批量动作统一遵循“先预览、再执行、可回退”的事务流程，"
    "先在只读模式下模拟结果并生成对照表，确认后才真正执行，执行前先把日志落盘，"
    "任何一次批量操作都可以通过回退中心一键恢复，且不受系统撤销与时间限制。"
    "在此底座之上，工具提供批量重命名、批量格式转换、内容级去重、抽取式摘要、"
    "多格式导出与中文排版、本地扫描分类六个模块，并配套统一的回退中心与设置体系。"
)
body(
    "技术上，作品采用 .NET 8 + WPF 手写 MVVM 实现，无第三方依赖；文档解析与生成、图像编解码、"
    "PDF 排版、SimHash 文本指纹与 TextRank/TF-IDF 摘要全部在本地纯代码完成，"
    "不调用任何外部 API、不加载任何 AI 模型。最终以自包含单文件发布，体积 68.5 MB，"
    "目标电脑双击即用，无需安装 .NET 运行环境，4 GB 内存设备可流畅运行。"
    "项目通过 17 组自动化自测、126 项断言验证核心链路，并在真实文件夹上完成端到端操作验证。"
)
body(
    "本作品的开发过程本身是 AIGC 应用实践：需求分析、方案设计、界面代码与大部分业务代码由 AI 编码工具"
    "在人工设计的分步 Prompt 下生成，团队负责需求判断、技术决策、逐项验证与缺陷修复。文档附录完整记录了"
    "AI 工具使用情况、关键 Prompt、生成与开发过程、迭代优化、测试验证和人工核查情况。"
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
r = p.add_run("关键词：")
run_font(r, HEI, 12, bold=True)
r = p.add_run("本地文件处理；可回退事务；SimHash 文本指纹；离线软件；AIGC 辅助开发")
run_font(r, SONG, 12)

# ================================================================ 目录
doc.add_section(WD_SECTION.NEW_PAGE)
heading("目　录", 1)

toc_p = doc.add_paragraph()
toc_p.paragraph_format.space_before = Pt(6)
run = toc_p.add_run()
run_font(run, SONG, 11)
fld_begin = OxmlElement("w:fldChar")
fld_begin.set(qn("w:fldCharType"), "begin")
fld_begin.set(qn("w:dirty"), "true")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = ' TOC \\o "1-2" \\h \\z \\u '
fld_sep = OxmlElement("w:fldChar")
fld_sep.set(qn("w:fldCharType"), "separate")
placeholder = OxmlElement("w:t")
placeholder.text = "（在 Word 中右键选择“更新域”，即可生成带页码的目录）"
fld_end = OxmlElement("w:fldChar")
fld_end.set(qn("w:fldCharType"), "end")
for el in (fld_begin, instr, fld_sep, placeholder, fld_end):
    run._r.append(el)

# ================================================================ 正文
doc.add_section(WD_SECTION.NEW_PAGE)
body_sec = doc.sections[-1]
pg = OxmlElement("w:pgNumType")
pg.set(qn("w:start"), "1")
body_sec._sectPr.append(pg)

# ---------------------------------------------------------------- 一
heading("一、项目背景", 1)

heading("1.1 问题来源", 2)
body(
    "学生和年轻办公人群每天都在和本地文件打交道：整理课程资料、转换文档格式、批量重命名、处理重复内容。"
    "这些事看起来简单，却反复消耗着大量时间——一次课程作业可能要处理几十份 PDF 与 PPT，"
    "一次社团活动可能要整理上百张照片，一次资料归档可能要在十几个文件夹之间反复搬运文件。"
)
body(
    "这类问题有三个特征：发生频率介于“每天”和“每月几次”之间，影响程度多为“中等”，"
    "单独看都不严重，但长期累积会持续消耗注意力；解决问题的门槛很低，用户往往选择手动硬扛；"
    "一旦出错代价很高，改错几百个文件名、误删一份资料，都需要花费成倍的时间补救。"
    "因此它不是“致命痛点”，而是典型的“烦人问题”——最适合用轻量、顺手、低门槛的工具解决，"
    "而不适合用重型平台。"
)

heading("1.2 用户调研：27 份有效样本的关键发现", 2)
body(
    "团队围绕本地工具软件的使用习惯、痛点与偏好设计问卷，回收有效样本 27 份，"
    "以学生和刚入职、实习的年轻用户为主，地域覆盖浙江、四川、湖南、辽宁、重庆等 11 个省市，"
    "全部以简体中文作答。调研的完整统计与图表见《本地工具软件用户需求调研分析报告》，"
    "本项目只摘录与产品决策直接相关的结论。"
)
caption("表 1　用户调研的关键结论摘要", "数据来源：本地工具软件用户需求调研（N=27，多选题为人数约数）")
add_table(
    ["调研维度", "关键结论", "数据支撑"],
    [
        ["使用频率", "是具有一定日常粘性的刚需场景", "约 67% 每周至少使用几次（每周几次 12 人、每天 6 人）"],
        ["使用场景", "学习辅助、办公事务、校园服务居前三", "三类场景高度重叠，用户需处理课程资料、PPT、格式转换与重复内容"],
        ["核心痛点", "文件整理耗时、格式转换麻烦、操作步骤太多", "三类问题选择人数最多，是最高频的痛点组合"],
        ["现有方案不满", "耗时、步骤多、结果不准确、需要付费", "耗时 15 人、步骤太多 12 人、结果不准确 10 人、需要付费 10 人"],
        ["尝试意愿", "尝试意愿非常高", "约 89% 愿意或非常愿意尝试（非常愿意 14 人、愿意 10 人）"],
        ["产品形式", "桌面软件是首选形态", "桌面软件 20 人，本地网页应用 12 人，命令行工具仅 5 人"],
        ["最看重的方面", "隐私安全 > 离线可用 > 易用性", "前三项排序稳定，免费仅排第 6 位"],
        ["主要顾虑", "安装复杂、占用资源大、功能效果差", "安装复杂 12 人、占资源 11 人、效果差 11 人、兼容性 9 人"],
    ],
    widths=[3.2, 5.4, 7.4],
)
body(
    "调研给出三条可以直接指导设计的信号。第一，用户要的不是“更多功能”，而是“更少步骤”和“更准的结果”，"
    "“省时、少步骤、结果准”构成产品价值主张的核心三角。第二，隐私安全与离线可用被排在最看重的前两位，"
    "但“不介意上传”的人极少，说明用户真正需要的是数据控制权，而不是对云服务的绝对排斥。"
    "第三，用户最大的心理门槛不是“有没有功能”，而是“装起来麻烦、跑起来卡、效果不好”，"
    "因此“双击即用、体积小、不占资源”本身就是功能。"
)

heading("1.3 现有工具的四个缺口", 2)
body(
    "把调研结论与现有工具的实际能力对照后，团队确认市场存在四个明确缺口。这些缺口不是因为技术上做不到，"
    "而是因为主流工具把注意力放在了功能数量上，忽略了“用得放心”这件事。"
)
dash_item(
    "现有批量重命名工具一旦执行，撤销功能往往已经失效，改错的几百个文件名只能靠记忆恢复。"
    "这种恐惧让很多用户宁可手动一个一个改，也不敢用批量工具。",
    bold_prefix="批量操作缺少可回退机制。",
)
dash_item(
    "佳能、尼康、索尼、富士、松下的 RAW 格式高度封闭，通用转换工具往往只支持部分格式，"
    "专业软件则需要付费订阅，摄影爱好者被迫在“付费”和“放弃”之间选择。",
    bold_prefix="跨厂商 RAW 格式支持不足。",
)
dash_item(
    "现有去重工具主要按文件名或字节哈希比对，文件改了名字或经过格式转换，工具就认不出来了，"
    "而这恰恰是资料重复最常见的形态。",
    bold_prefix="内容级去重缺位。",
)
dash_item(
    "大量轻量工具优先考虑英文场景，中文排版、标题层级、上下标、模板样式都被牺牲，"
    "导出后往往需要手工重排一遍，等于白做。",
    bold_prefix="中文场景的导出质量差。",
)

heading("1.4 机会判断", 2)
body(
    "与此同时，用户对本地工具的偏好正在变强：隐私安全、离线可用和易用性被视为本地工具的核心卖点，"
    "“完全本地处理”与“可联网但数据不上传”合计获得样本中的绝大多数选择。"
    "技术条件也已经成熟：.NET 自包含单文件发布解决了“不装环境”的问题，"
    "WIC 图像编解码与纯代码的 OOXML 解析足以覆盖常见文档与图片格式，"
    "文本指纹与图排序算法可以在毫秒级完成相似度计算。"
)
body(
    "据此可以给出一个明确判断：做一个完全离线的本地文件处理工具，在“省时、少步骤、结果准”上做到位，"
    "同时在“可回退”这个没人愿意做的环节上做扎实，就具备真实的市场空隙。"
)

# ---------------------------------------------------------------- 二
heading("二、解决方案", 1)

heading("2.1 产品定位与设计原则", 2)
body(
    "LocalCraft 是一款完全离线的本地文件处理工具，以 Windows 桌面软件（exe）为主要形态。"
    "核心定位是：不联网、不装环境、不占内存、U 盘即插即用。所有文件处理都在本机完成，数据不出本机。"
)
body("围绕这一定位，作品确定了五条贯穿始终的原则：", indent=True)
dash_item("所有批量操作必须遵循“先预览、再执行、可回退”的流程，没有例外。", bold_prefix="① 安全优先。")
dash_item("核心功能完全离线可用，程序不包含 HTTP 客户端、遥测与更新检查。", bold_prefix="② 离线可用。")
dash_item("自包含单文件发布，目标电脑不需要安装 .NET 或任何运行库。", bold_prefix="③ 零依赖分发。")
dash_item("体积控制在几十 MB，启动快、内存占用低，4 GB 内存设备可流畅运行。", bold_prefix="④ 轻量。")
dash_item("不做功能堆砌，把用户最容易被忽视、却最影响使用信心的环节做扎实。", bold_prefix="⑤ 克制。")

heading("2.2 功能框架", 2)
body(
    "作品以“可回退的批量操作”为底座，向上生长出六个业务模块，并由回退中心与设置页统一支撑。"
    "六个模块共享同一套日志格式、预览交互与回退逻辑，因此用户学一次即可迁移到其它模块。"
)
caption("表 2　LocalCraft 的功能模块划分", "“状态”列为 v0.4 的实际交付状态")
add_table(
    ["模块", "核心能力", "状态"],
    [
        ["模块 1　批量重命名", "前缀/后缀/查找替换/序号，只读预览、冲突检测、覆盖前备份、一键回退", "可用"],
        ["模块 2　格式转换", "文档与图片批量互转，队列、进度、取消清理，输出不覆盖已有文件", "可用"],
        ["模块 3　内容级去重", "文本提取 + SimHash 指纹 + 汉明距离分组，按阈值展示相似组并隔离重复项", "可用"],
        ["模块 4　抽取式摘要", "TextRank 关键句 + TF-IDF 关键词，附带原文行号与字符偏移", "可用"],
        ["模块 5　导出与排版", "Markdown / Word DOCX / PDF 多格式导出，中文标题层级与版式", "可用"],
        ["模块 6　扫描与分类", "只读扫描，按类型/年份/年月/名称前缀预览并移动，可回退", "可用"],
        ["回退中心", "统一读取本机日志，恢复重命名与移动，删除转换与导出输出", "可用"],
        ["设置与常用设置", "动效等级、默认值、页面记忆、快捷设置面板", "可用"],
    ],
    widths=[4.0, 9.4, 2.2],
)

heading("2.3 统一安全底座：先预览、再执行、可回退", 2)
body(
    "所有批量动作遵循同一套流程，这套流程是本作品与同类工具最本质的差别："
)
body(
    "第一步，只读规划。工具按用户规则在内存中计算每个文件的“旧状态 → 新状态”，"
    "不创建目录、不写入任何文件，预览阶段零写入。第二步，对照确认。"
    "对照表逐项标明每个文件是“可执行”“覆盖”“冲突”“无变化”还是“无效”，"
    "冲突项会直接阻断提交，直到用户调整规则或取消勾选。"
    "第三步，日志先行。用户确认后，工具先生成并落盘操作日志（状态为“已准备”），再开始动手，"
    "确保任何中断都留得下证据。第四步，事务执行与回退。执行过程中逐项更新状态，失败则逆序回滚；"
    "全部完成后日志标记为“已完成”，随时可以在回退中心一键恢复。"
)
caption("表 3　批量操作中单个文件的状态定义")
add_table(
    ["状态", "含义", "界面处理"],
    [
        ["可执行", "目标名称合法且不存在同名文件", "默认勾选，参与执行"],
        ["覆盖", "目标已存在同名文件；覆盖前先把原文件移入 backups 目录", "需用户勾选“允许覆盖”才执行"],
        ["冲突", "批次内多个文件指向同一目标，或目标被批次内其它源文件占用", "阻断提交，需调整规则"],
        ["无变化", "规则作用后文件名没有变化", "默认不勾选"],
        ["无效", "空文件名、非法字符、Windows 保留名、路径过长等", "阻断提交，需修正规则"],
    ],
    widths=[2.6, 8.4, 4.6],
)
body(
    "之所以把“可回退”做成底层能力而不是附加功能，是因为它直接决定用户“敢不敢用”。"
    "同类工具的撤销依赖系统临时状态，窗口一关就失效；而本作品把操作前后的完整映射写进本机日志文件，"
    "回退时逆向执行，既不依赖系统撤销，也不受时间限制。"
)

heading("2.4 六大功能模块", 2)

heading("2.4.1 模块 1：批量重命名", 3)
body(
    "这是全套机制的原型模块。用户选择工作文件夹与扩展名筛选后，可组合前缀、后缀、查找替换、"
    "序号（起始值、位数、分隔符、是否前置）等规则，实时生成预览对照表，"
    "并给出“待重命名 / 覆盖 / 冲突 / 无变化”的计数。执行前可勾选是否允许覆盖；"
    "执行完成后，回退中心可以一键把文件名恢复原状，被覆盖的文件也会从备份中还原。"
)

heading("2.4.2 模块 2：批量格式转换", 3)
body(
    "文档支持 DOCX、PPTX、Markdown、TXT、HTML 输入，输出 Markdown、TXT、HTML、DOCX、PDF；"
    "图片支持 JPG、PNG、BMP、GIF、TIF 互转，使用 Windows 自带的 WIC 解码器，"
    "透明通道转 JPEG 时自动以白底合成。文章转出时不覆盖已有文件，同名自动追加序号；"
    "转换过程显示队列与进度，可随时取消并清理临时输出。"
)
body(
    "转换的回退逻辑与重命名不同：原始文件从未被修改，回退只需删除该批次生成的输出文件，"
    "因此“转换失败”对用户是零损失。"
)

heading("2.4.3 模块 3：内容级去重", 3)
body(
    "用户添加文档或文件夹后，工具提取正文（DOCX 读段落与标题样式，PPTX 按页读取，"
    "PDF 读取文本层），中文切成字符二元组、英文切成词元，再生成 64 位 SimHash 指纹，"
    "按汉明距离把相似文档聚成连通组。界面按相似组展示基准文件、相似度与隔离目标位置，"
    "相似度阈值可在 50%～100% 之间调节。确认后，勾选的重复项被移动到隔离文件夹，"
    "移动过程写入日志，可一键恢复。"
)
figure("module3-dedup.png")
caption("图 1　内容级去重界面：相似组、阈值调节与隔离策略", "图片来源：LocalCraft v0.4 实际运行截图")

heading("2.4.4 模块 4：抽取式摘要", 3)
body(
    "对单篇文档切句后，用句子相似度构建 TextRank 图并迭代得到关键句权重，"
    "同时用 TF-IDF 在句级文档频率上提取关键词。摘要结果只返回原文中的句子，不生成新内容，"
    "每一句都附带原文行号与字符偏移，可直接跳回原文核对。整个过程零配置、秒级返回，不加载任何模型。"
)
figure("module4-summary.png")
caption("图 2　抽取式摘要界面：关键句排序与原文位置追溯")

heading("2.4.5 模块 5：导出与中文排版", 3)
body(
    "输入 Markdown、TXT、HTML、DOCX、PPTX，可同时勾选 Markdown、Word DOCX、PDF 三种输出，"
    "一次生成多个结果。Markdown 与 DOCX 保留标题层级；PDF 由本地排版引擎生成："
    "A4 页面、微软雅黑字体、标题 1～3 级样式、中文段落间距、项目符号、粗体与基础上下标，"
    "版式稳定且不依赖外部字体解析。需要说明的是，PDF 页面以高分辨率位图封装，"
    "中文显示可靠，但文字在阅读器中不可选中，这是当前版本明确的取舍。"
)
figure("module5-export.png")
caption("图 3　导出与排版界面：多格式勾选、预览与统一日志")

heading("2.4.6 模块 6：扫描与分类", 3)
body(
    "扫描阶段严格只读：工具递归读取文件夹，按“文件类型 / 修改年份 / 修改年月 / 名称前缀”"
    "四类规则计算每个文件应去的分类子文件夹，并单独标记冲突项与已归类文件。"
    "用户确认后才创建子文件夹并移动文件，移动复用批量执行器，日志记录完整映射，"
    "回退中心可把文件逐一移回原位置。"
)

heading("2.4.7 回退中心", 3)
body(
    "回退中心集中读取 %LOCALAPPDATA%\\LocalCraft\\journals 下的日志文件，按时间列出历史批次。"
    "对重命名与移动，它逆向恢复文件路径并还原覆盖备份；对转换与导出，它只删除该批次生成的输出，"
    "原始文件不受影响。每条日志独立存在，只要文件还在，回退随时可用。"
)

heading("2.4.8 设置与常用设置", 3)
body(
    "顶部栏齿轮按钮打开“常用设置”弹出面板，可直接调整动画开关、页面记忆、默认包含子文件夹、"
    "去重阈值；完整设置页提供动效等级（轻量 / 标准 / 丰富）与动画预览、摘要关键句数量、"
    "默认分类方式、数据目录入口与恢复默认设置。设置原子写入本机 settings.json，不涉及任何云端同步。"
)
figure("module7-settings.png")
caption("图 4　设置页：动效等级、默认值与本地数据目录")

heading("2.5 界面与交互", 2)
body(
    "界面采用左侧导航 + 右侧工作的两栏结构。左侧导航固定八个入口（六个业务模块、回退中心、设置），"
    "每项都配有矢量线性图标与一句能力说明；顶部状态条常驻“离线可用”标识，"
    "让用户随时确认工具没有联网。工作区顶部是模块标题与一句话说明（例如“先预览、再执行、可回退。"
    "预览阶段只读，不修改任何文件。”），中部是参数区，下部是预览表格，"
    "底部固定执行按钮与安全提示。凡是会改动文件的操作，按钮旁都有一行说明写清风险与恢复方式。"
)
figure("module1-home.png")
caption("图 5　主界面：批量重命名模块的预览态（尚未执行任何写操作）")

# ---------------------------------------------------------------- 三
heading("三、技术实现路径", 1)

heading("3.1 技术选型", 2)
body(
    "选型的唯一标准是“能否在完全离线、不装环境、轻量的前提下把功能做对”。"
    "团队放弃了 Web 技术栈（需要内嵌浏览器内核，体积与内存不达标）、"
    "放弃了 Python + 打包器的方案（体积偏大、启动偏慢），最终选择 .NET 8 + WPF。"
)
caption("表 4　关键技术选型")
add_table(
    ["主题", "选择", "选择理由"],
    [
        [".NET 8 + WPF", "自包含单文件 win-x64 发布", "目标机无需安装运行库；double buffer 渲染足以支撑千行表格"],
        ["界面架构", "手写 MVVM，无第三方 NuGet 包", "无供应链风险，发布体积可控，逻辑可测"],
        ["文档处理", "自研 OOXML / Markdown / HTML 解析与写入", "不依赖 Office 与外部引擎，离线可用"],
        ["图像处理", "Windows Imaging Component（WIC）", "系统自带解码器，覆盖 JPG/PNG/BMP/GIF/TIF"],
        ["PDF 输出", "WPF FlowDocument 排版 → 页面位图 → 本地 PDF 封装", "中文渲染可靠，无需外部字体解析"],
        ["去重", "中文 bigram / 英文词元 + 64 位 SimHash", "对改写、改名、格式转换鲁棒"],
        ["摘要", "TextRank 句子图 + TF-IDF", "无模型、秒级、结果可追溯"],
        ["日志", "System.Text.Json，原子写入", "先落盘再动手，断电也留得下证据"],
    ],
    widths=[3.2, 6.4, 5.9],
)

heading("3.2 系统分层与项目结构", 2)
body(
    "系统分为三层：UI 层（LocalCraft.App）、转换内核（LocalCraft.Conversion）与核心层（LocalCraft.Core）。"
    "核心层不引用 WPF，因此可以在控制台里直接测试；转换内核只依赖 Windows 的图像与排版能力；"
    "UI 层通过普通 C# 对象与下层交互，文件系统逻辑不写在 code-behind 里。"
)
caption("表 5　三个程序集的职责划分")
add_table(
    ["程序集", "职责", "关键类型"],
    [
        ["LocalCraft.Core", "重命名规划与冲突检测、批量事务与回滚、JSON 日志与回退、SimHash/TextRank/TF-IDF、分类规划、设置持久化", "RenamePlanner、BatchRenameExecutor、UndoService、JournalRepository、DedupEngine、TextSummaryEngine、ClassifyPlanner、AppSettingsService"],
        ["LocalCraft.Conversion", "输出路径规划、引擎匹配、文档/图像/PDF 转换、外部引擎适配", "ConversionPlanner、ConversionExecutor、BuiltInDocumentEngine、BuiltInImageEngine、PdfRenderer、ConversionEngineCatalog"],
        ["LocalCraft.App", "WPF 界面、手写 MVVM、导航与动效、设置页与常用设置面板", "MainViewModel、RenameViewModel、ConvertViewModel、DedupViewModel、SummaryViewModel、ExportViewModel、ClassifyViewModel、UndoCenterViewModel、SettingsViewModel"],
    ],
    widths=[3.4, 7.2, 4.9],
)

heading("3.3 关键实现一：只读规划与事务执行", 2)
body(
    "重命名规划器 CreatePlan() 只计算不写盘。输入文件先规范化、去重、自然排序，然后逐项应用规则，"
    "并校验文件名（空主体、非法字符、结尾空格或句点、Windows 保留名、组件长度、完整路径长度），"
    "再检测目标文件夹、目标同名文件、批次内目标重复、批次内源文件占用；"
    "仅大小写改变的文件标为可执行，由执行器用临时文件中转。"
)
body(
    "执行阶段的顺序是设计重点：先再次预检，防止预览后文件系统发生变化；"
    "然后生成状态为“已准备”的日志并落盘；若目标已存在且允许覆盖，先把目标移入 backups\\<批次号>\\ 并记录备份路径；"
    "再把源文件重命名到目标，周期性落盘（每 25 个操作一个检查点）；"
    "任一步失败则逆序恢复已完成的改名并还原备份，日志标记为“已回滚”。"
)
body(
    "回退是同一套机制的逆过程：重新加载日志，校验状态与操作完整性；预检目标文件、原路径占用与备份是否存在；"
    "把日志标记为“回退中”；逆序执行“新文件 → 原路径”，覆盖过的目标从备份还原；"
    "最后标记“已回退”并清理该批次备份。中途失败会尝试补偿，把文件恢复到回退前的状态。"
)

heading("3.4 关键实现二：本地转换与导出内核", 2)
body(
    "文档引擎用 System.IO.Compression 与 XDocument 直接读写 Open XML 包："
    "读 DOCX 时解析 word/document.xml 的段落、标题样式、列表与表格文本，读 PPTX 时按 slideN.xml 逐页取文本；"
    "写 DOCX 时生成最小 Open XML 包（[Content_Types].xml、_rels/.rels、word/document.xml、styles.xml），"
    "内置微软雅黑默认字体、Normal 与 Heading1～3 样式、A4 页面与段落间距。整个过程不需要安装 Office。"
)
body(
    "图像引擎基于 WIC 顺序编解码：每张图在内存中处理完立即写盘，批次不在内存里累积全部图片，"
    "因此三千张图片的任务也不会把内存打满。PDF 输出走另一条路径：构造 A4 的 FlowDocument，"
    "经 DocumentPaginator 分页，以 144 DPI 把每页渲染成位图并用 Flate 压缩，"
    "最后封装成一个最小 PDF，每页使用一个图像 XObject。"
)
body(
    "外部引擎是可选增强而非依赖：程序启动时探测本机是否存在 Pandoc 与 LibreOffice，"
    "只有在检测到对应引擎时，PDF 输入、旧版 DOC/PPT 与 PPTX 输出才会进入执行队列；"
    "调用时通过参数列表传递，不使用 shell 拼接，并设置超时与进程树清理。"
)

heading("3.5 关键实现三：内容指纹与摘要算法", 2)
body(
    "文本提取器把 DOCX、PPTX、Markdown、TXT、HTML 与 PDF 文本层统一转成纯文本；"
    "分词器把中文切成字符二元组、英文切成词元，对 64 个位分别加权投票，得到 64 位 SimHash 指纹。"
    "去重时以汉明距离衡量相似度，并用并查集做连通分组，使“A 像 B、B 像 C”的链条能被识别为同一组。"
    "这样，改过文件名、换过格式、重排过段落的同一份资料依然会被聚到一起。"
)
body(
    "摘要算法按中文句号、问号、感叹号、分号与换行切句，保留字符偏移与行号；"
    "句子两两相似度构成 TextRank 图，迭代求得句子权重；TF-IDF 在句级文档频率上提取关键词。"
    "输出严格取原文句子，保证“可追溯”而不是“看起来像”。"
)

heading("3.6 工程化：性能、兼容与隐私", 2)
caption("表 6　工程指标与实现方式")
add_table(
    ["指标", "实际情况", "实现方式"],
    [
        ["发布体积", "68.5 MB 单文件", "自包含 win-x64 裁剪发布，不含浏览器内核与模型文件"],
        ["运行环境", "双击即用，无需安装 .NET", "自包含单文件，目标机不需要任何运行库"],
        ["内存占用", "4 GB 内存设备可流畅运行", "表格行列虚拟化、图像顺序处理、无后台常驻服务"],
        ["单次处理规模", "重命名 10 000 / 转换与分类 5 000 / 去重 500 篇", "分批读取 + 流式写盘，界面只渲染可见行"],
        ["高 DPI", "100%～200% 缩放不横向裁切", "Per-Monitor V2 + 1320×820 设计画布 + Viewbox 等比适配"],
        ["隐私", "无网络调用", "不引用 HttpClient、WebView、遥测 SDK；数据目录可用环境变量重定向"],
        ["数据留存", "日志与备份只在本机", "%LOCALAPPDATA%\\LocalCraft\\{journals,backups}"],
    ],
    widths=[3.0, 6.2, 6.3],
)

heading("3.7 开发环境与发布流程", 2)
body(
    "开发环境为 Windows 10/11 + .NET 8 SDK，构建、自测与发布均由命令行脚本完成，"
    "便于在资源受限的机器上复现（构建时使用 -m:1 降低并发）。"
)
prompt_block([
    "# 构建",
    "dotnet build .\\LocalCraft.sln -c Release -m:1",
    "",
    "# 运行",
    "dotnet run --project .\\src\\LocalCraft.App\\LocalCraft.App.csproj -c Release",
    "",
    "# 核心自测",
    "dotnet run --project .\\tests\\LocalCraft.SelfTest\\LocalCraft.SelfTest.csproj -c Release",
    "",
    "# 发布自包含单文件",
    "powershell -ExecutionPolicy Bypass -File .\\scripts\\publish-win-x64.ps1",
], indent=0.5)

# ---------------------------------------------------------------- 四
heading("四、团队分工", 1)
body(
    "团队按“产品—设计—算法—工程—验证”五条线分工，每人都要参与需求评审与交叉验收，"
    "避免出现“只写代码不理解需求”或“只做设计不碰实现”的断点。"
    "由于开发主要由 AI 编码工具承担，团队的核心工作从“写代码”转向“定义问题、拆解任务、验证结果”。"
)
caption("表 7　团队分工")
add_table(
    ["角色", "姓名", "主要职责", "对应模块 / 产出"],
    [
        ["队长 / 产品", "　　", "需求调研设计与分析、产品定位、功能优先级、验收标准制定、对外汇报", "调研报告、PRD、答辩稿"],
        ["交互与视觉", "　　", "页面结构、导航与动效、图标矢量绘制、高 DPI 适配、界面截图与演示脚本", "界面原型、主题与样式、演示视频"],
        ["算法", "　　", "SimHash 指纹与分组、TextRank/TF-IDF 摘要、文本提取、相似度阈值调参", "模块 3、模块 4"],
        ["工程 / 转换", "　　", "OOXML 与 Markdown 读写、WIC 图像、PDF 排版、外部引擎探测、发布脚本", "模块 2、模块 5、发布包"],
        ["测试 / 文档", "　　", "自动化自测用例、边界与异常验证、文档排版、AI 生成内容核查", "自测套件、项目文档、核查记录"],
    ],
    widths=[2.8, 2.0, 6.6, 4.1],
)
body(
    "协作机制上，团队采用“任务卡片 + Prompt 模板 + 验收清单”的方式推进："
    "每完成一个功能点，先由负责人写出可执行的需求描述与验收标准，再整理成结构化 Prompt 交给 AI，"
    "AI 产出后由负责人逐项运行验证，通过后合并并更新文档。所有 Prompt 与修改记录都保留在仓库中，"
    "因此任何一个模块都能追溯到“为什么这么做”。"
)

# ---------------------------------------------------------------- 五
heading("五、创新点与价值", 1)

heading("5.1 创新点一：把“可回退”做成批量操作的底层能力", 2)
body(
    "现有批量工具的通行做法是“执行完就结束”，撤销依赖系统临时状态，窗口一关就失效。"
    "本作品把“先预览、再执行、可回退”作为所有批量动作的统一流程，用日志文件记录操作前后的完整映射，"
    "回退不依赖系统、不受时间限制，且覆盖重命名、移动、转换、导出四类动作。"
    "这解决的是用户最深层的心理障碍——不是“能不能做”，而是“敢不敢做”。"
    "价值在于：它把批量操作从“有风险的高效工具”变成“可以放心用的日常工具”。"
)

heading("5.2 创新点二：用内容相似度去重，替代文件名与字节比对", 2)
body(
    "现有工具只能识别逐字节完全相同的文件，文件改名或经过格式转换后就失效。"
    "本作品对文档提取文本后用 SimHash 计算相似度，并以连通分组识别“内容高度相似但格式或文件名不同”的重复文档，"
    "相似度阈值交给用户决定，隔离动作同样纳入回退体系。"
    "价值在于：它匹配的是用户真实的认知——“这两份资料是不是同一份东西”，"
    "而不是机器的认知——“这两个文件是不是同一串字节”。"
)

heading("5.3 创新点三：在完全离线的条件下补齐 RAW 与中文排版短板", 2)
body(
    "RAW 格式长期依赖付费专业软件，中文 Markdown 与 PDF 导出在轻量工具中普遍翻车。"
    "本作品通过系统 WIC 解码器尝试 RAW/HEIC 输入，并以自研的本地中文排版引擎生成 PDF，"
    "在纯本地、零付费的前提下覆盖中文标题层级、段落间距与基础上下标。"
    "价值在于：它服务的是被主流工具忽视的具体人群——摄影爱好者与中文用户；"
    "同时也如实说明当前边界：LibRaw 尚未随包分发，缺少解码器时任务会明确失败并清理临时文件，而不是静默出错。"
)

heading("5.4 与现有方案的对比", 2)
caption("表 8　LocalCraft 与现有方案的对比")
add_table(
    ["对比维度", "常见批量改名工具", "在线转换网站", "通用 AI 助手", "LocalCraft"],
    [
        ["可回退", "依赖系统撤销，窗口关闭即失效", "不涉及", "需用户自己描述回滚方案", "日志驱动，任意时间可回退"],
        ["去重依据", "文件名 / 字节哈希", "不支持", "需上传后才能比较", "文本指纹 + 相似度分组"],
        ["离线与隐私", "多数本地，但功能零散", "必须上传文件", "必须联网", "全本地，无网络调用"],
        ["中文导出", "一般不涉及", "版式常丢", "需往返复制粘贴", "本地排版，层级与样式保留"],
        ["安装与体积", "需安装，体积不定", "无需安装", "无需安装", "单文件 68.5 MB，无需安装"],
        ["费用", "免费或订阅", "免费额度 + 付费", "订阅制", "本地运算，零边际成本"],
    ],
    widths=[2.6, 3.4, 3.0, 2.9, 3.6],
)

heading("5.5 应用价值", 2)
body(
    "对个人用户，它把“整理资料”从一件要下决心的事，变成一件顺手就能做的事；"
    "每一次操作都可以撤销，意味着用户可以放心尝试规则，学习成本随之下降。"
    "对校园与小型办公场景，它在不引入任何云服务的前提下提供了可审计的文件整理能力，"
    "日志文件本身就是操作留痕，便于事后复核。"
)
body(
    "从 AIGC 应用的角度看，本作品验证了一条可行的路径："
    "把“需求定义、方案拆解、验收标准”留给团队，把“代码实现、界面搭建、文档整理”交给 AI，"
    "用可运行的原型和自动化自测来约束生成质量。这条路径对小团队尤其有效。"
)

# ---------------------------------------------------------------- 六
heading("六、总结与展望", 1)

heading("6.1 项目总结", 2)
body(
    "本项目从学生和年轻办公人群的真实文件处理痛点出发，识别出市场在批量操作安全性、RAW 格式支持、"
    "内容级去重、中文导出质量四个方向的缺口，提出并实现了一个完全离线的本地文件处理工具。"
)
body(
    "方案以“可回退的批量操作”为核心底座，覆盖批量重命名、格式转换、内容级去重、抽取式摘要、"
    "中文优化导出、本地扫描分类六个模块，全部在本地完成，不联网、不装环境、不占内存。"
    "技术上用 .NET 8 + WPF 自包含单文件落地，体积 68.5 MB；核心链路经 17 组自测、126 项断言验证，"
    "在真实文件夹上完成了端到端操作验证。"
)
body(
    "更重要的是，团队在过程中形成了可复用的 AI 协作方法：把模糊需求拆成可验收的任务，"
    "用结构化 Prompt 约束 AI 的输出形态，再用自动化测试把“看起来能用”变成“确实可用”。"
)

heading("6.2 当前边界", 2)
body("作品在完成度的同时仍有明确边界，团队选择如实说明而不是回避：")
caption("表 9　当前版本的已知边界与应对")
add_table(
    ["边界", "现状", "应对方式"],
    [
        ["RAW / HEIC 输入", "LibRaw 未随包分发，只尝试系统 WIC 解码器", "检测不到时明确失败并清理临时输出，界面给出原因"],
        ["PDF 输入、旧版 DOC/PPT、PPTX 输出", "依赖本机 LibreOffice", "启动时探测引擎，缺失时任务不进入队列"],
        ["PDF 导出文字不可选中", "页面以 144 DPI 位图封装", "版式优先；后续计划增加文本层"],
        ["扫描版 PDF", "无文本层，不做 OCR", "在界面说明跳过原因"],
        ["复杂 Word 版式", "文档转换以文本级内容为主", "不承诺保留分栏、公式与图片版式"],
        ["中断事务", "断电或强杀造成的未完成批次不自动恢复", "日志保留供人工核对，正常完成批次可长期回退"],
    ],
    widths=[3.6, 6.2, 5.7],
)

heading("6.3 后续展望", 2)
body("工具稳定后，团队计划沿三个方向继续推进：")
dash_item("打包 LibRaw 与 LibreOffice 可选引擎，把 RAW、旧版 Office 与 PDF 输入做成完全离线的完整闭环。", bold_prefix="① 格式覆盖。")
dash_item("为分类增加“类型 + 年份”等规则组合与空目录清理，让整理更贴合个人习惯。", bold_prefix="② 规则引擎。")
dash_item("为 PDF 增加文本层、为去重组提供任意基准文件选择与回收站选项；探索可选的增强能力，但始终不牺牲隐私与可控性。", bold_prefix="③ 体验增强。")
body(
    "最终目标不是做一个功能最多的工具，而是做一个用户敢用、愿用、用得放心的本地文件处理工具。"
)

# ---------------------------------------------------------------- 附录
page_break()
heading("附录 A　AI 工具使用情况说明", 1)
body(
    "按照大赛“作品中如使用 AIGC 工具生成内容，须在项目文档中明确说明”的要求，"
    "本节如实记录本项目使用过的 AI 工具、使用环节与人工处理方式。"
)

heading("A.1 使用的 AI 工具与环节", 2)
caption("表 A1　AI 工具使用清单")
add_table(
    ["AI 工具", "使用环节", "具体用途", "人工处理"],
    [
        ["豆包（文档助手）", "需求与文档", "梳理调研结论、起草项目文档与汇报材料、调整文档格式与措辞", "团队核对数据出处、重写关键判断、统一术语"],
        ["DeepSeek（通用大模型）", "需求分析、方案设计、排错", "拆解功能、讨论技术路线、解释报错原因、评审设计取舍", "逐条评估可行性，删除不符合离线约束的建议"],
        ["DeepSeek Harness（AI 编码智能体）", "代码实现与调试", "生成本项目绝大部分 C# / XAML 代码，执行构建、运行自测、修复缺陷、更新文档", "定义任务与验收标准，逐项运行验证，人工重写不合理实现"],
    ],
    widths=[3.6, 3.2, 5.6, 3.1],
)

heading("A.2 生成内容声明", 2)
body(
    "界面布局、图标矢量数据、绝大部分业务逻辑代码、本文件的部分措辞由 AI 生成并经团队修改。"
    "产品定位、四个市场缺口的判断、技术选型（.NET 8 + WPF 而非 Python/Web 方案）、"
    "可回退机制的流程设计、模块边界与已知边界的取舍，由团队讨论确定，AI 只参与实现与讨论。"
)
body(
    "所有 AI 生成内容均经过人工核查：代码必须通过编译与自测，界面必须在 100%～200% 缩放下实际运行验证，"
    "文档中的每一个数据必须能追溯到调研报告或自测日志。项目未使用他人的代码、素材或数据，"
    "未使用任何在线 API 处理用户文件，作品本体不包含 AI 模型。"
)

heading("附录 B　关键 Prompt", 1)
body(
    "本附录保留实际使用过的关键 Prompt。团队的 Prompt 写法遵循四条原则："
    "先给角色与目标，再给硬性约束；把验收标准写进指令，让“完成”可判定；"
    "把大任务拆成可交付的小步骤，逐步验证；要求直接给出可运行结果，而不是思路。"
)

heading("B.1 核心开发 Prompt（节选）", 2)
body("这是项目的起点 Prompt，决定了整个产品的边界与形态：", indent=True)
prompt_block([
    "请帮我开发一个名为 LocalCraft 的 Windows 桌面软件。",
    "",
    "一、产品定位",
    "一款完全离线的本地文件处理工具。不联网、不装环境、不占内存、U 盘即插即用。",
    "所有文件处理都在本机完成，数据不出本机。",
    "",
    "二、核心原则",
    "1. 所有批量操作必须遵循“先预览、再执行、可回退”的流程。",
    "2. 核心功能完全离线可用，不依赖任何网络请求。",
    "3. 打包为自包含 exe，目标电脑不需要安装任何运行库。",
    "4. 轻量优先，体积控制在几十 MB 到一百多 MB，4GB 内存可流畅运行。",
    "",
    "三、功能模块",
    "模块 1：可回退的批量操作（核心底座）",
    "　用户设定规则后，先在只读模式下模拟操作结果，输出“旧状态 → 新状态”的预览对照表；",
    "　对照表需标明：哪些文件会被重命名、哪些会被覆盖、哪些存在冲突；",
    "　用户确认后才真正执行；执行时自动保存可回滚的日志文件，记录操作前后的完整映射；",
    "　提供一键回退功能，回退不依赖系统撤销，不受时间限制；",
    "　该机制覆盖批量重命名、批量格式转换、批量移动与分类。",
    "模块 2：批量格式转换",
    "　文档：Word、PPT、PDF、Markdown、TXT 等常见格式互转；图片：JPG、PNG、HEIC 等；",
    "　RAW：支持主流厂商 RAW 格式；支持批量队列与进度显示；结果纳入可回退日志体系。",
    "模块 3：内容级去重",
    "　对 Word、PDF、TXT 等文档提取文本，使用 SimHash 或文本指纹计算相似度；",
    "　识别“内容高度相似但格式或文件名不同”的重复文档；",
    "　用户可设定相似度阈值，按阈值分组展示疑似重复项；处理动作纳入可回退体系。",
    "模块 4：抽取式文档摘要",
    "　不依赖 AI，采用 TextRank 或 TF-IDF 算法；提取关键句与关键词，按权重排序输出；",
    "　摘要结果可定位到原文位置，保证可追溯；零配置、秒出结果、不占内存。",
    "模块 5：导出与中文优化",
    "　支持导出为 Markdown、Word、PDF；重点保证中文场景质量：标题层级正确、中文排版不乱、",
    "　上下标与基础样式保留。",
    "模块 6：本地文件夹扫描与分类",
    "　按规则自动扫描文件夹，支持按类型、时间、名称等条件分类整理；",
    "　扫描阶段只读，不修改任何文件；分类动作纳入可回退的批量操作体系。",
    "",
    "四、技术约束",
    "开发语言：C#（.NET 8 + WPF）。",
    "去重算法：文本提取 + SimHash，纯本地计算。摘要算法：TextRank / TF-IDF，纯本地计算。",
    "回退机制：日志文件记录操作前后完整映射，回退时逆向执行。",
    "不使用任何 AI 模型，不接入任何外部 API。",
    "",
    "五、界面要求",
    "简洁、易用，降低学习成本；批量操作前必须展示预览对照表；",
    "提供清晰的回退入口；转换与去重过程显示进度。",
    "",
    "六、验收标准",
    "Windows 电脑上双击 exe 即可运行；断网状态下所有核心功能正常可用；",
    "批量重命名执行后，可通过日志一键回退到操作前状态；",
    "能识别内容相似但文件名不同的文档；中文 Markdown 导出后标题层级与基础样式正确；",
    "4GB 内存电脑可流畅运行。",
    "",
    "请先给出项目结构、模块划分和关键技术选型，再逐步实现各模块。从模块 1 开始实现，给出完整可运行代码。",
])

heading("B.2 迭代优化 Prompt", 2)
body("模块跑通后，界面与体验问题通过两条短指令完成迭代：", indent=True)
prompt_block([
    "① 先全部做完，并将左侧栏的图标换成图形式的。",
    "",
    "② 优化页面，加入设置及常用设置功能；加入更多动画效果。",
], indent=0.6)
body(
    "这两条指令的效果差异说明了一个经验：短指令只有在“上下文完整”时才有效。"
    "第①条之所以能直接落地，是因为前序对话已经建立了模块清单与界面结构；"
    "第②条之后，团队又补充了“动效必须可关闭”“设置写入本地 JSON”“高 DPI 下不能裁切”三条约束，"
    "否则 AI 会给出不符合离线与轻量要求的实现。"
)

heading("B.3 文档与工具类 Prompt（节选）", 2)
body("项目文档与辅助脚本也由 AI 协作完成，使用过的典型指令包括：", indent=True)
prompt_block([
    "① 根据 27 份问卷的统计结果，写一份用户需求调研分析报告，",
    "　 包含样本情况、核心发现（每项配数据与结论）、总结与建议，用 Word 可交付的格式。",
    "",
    "② 请帮我编写一个将本地文件/文件夹上传到 GitHub 仓库的脚本：使用 Git 命令行、",
    "　 支持指定分支与 commit message、自动 init/pull/add/commit/push、",
    "　 对 push 失败给出明确提示，并要求直接给出最终脚本代码与 2~3 个使用示例。",
], indent=0.6)

heading("附录 C　生成与开发过程", 1)
body(
    "项目按“调研 → 设计 → 底座 → 模块 → 体验 → 验证”的顺序推进。"
    "由于开发主要由 AI 编码智能体承担，团队把每个阶段的产出都定义为“可运行、可验证”的状态，"
    "避免出现大量无法验证的半成品代码。"
)
caption("表 C1　开发阶段与产出")
add_table(
    ["阶段", "主要工作", "AI 参与方式", "产出"],
    [
        ["第 5—6 周", "问卷设计、回收与分析，确认四个缺口，确定产品定位", "AI 辅助统计口径与报告结构，团队判断结论", "调研报告、产品定位"],
        ["第 7 周", "功能框架、交互草图、可回退流程与日志模型设计", "AI 参与方案讨论与取舍评估", "模块清单、流程图、日志字段设计"],
        ["第 8 周", "落地核心底座：规划器、执行器、日志与回退中心、批量重命名", "AI 生成代码，团队定义验收标准并逐项验证", "模块 1 与回退中心可用"],
        ["第 9 周", "转换内核：文档引擎、WIC 图像、PDF 排版；模块 2 与模块 5", "AI 生成代码与测试，团队验证中文版式", "模块 2、模块 5 可用"],
        ["第 10 周", "文本指纹与摘要算法，模块 3、模块 4、模块 6", "AI 生成算法实现，团队调参与边界测试", "模块 3、4、6 可用"],
        ["第 11 周", "界面重构、矢量图标、设置页与动效、性能与高 DPI 适配", "AI 生成 XAML 与动效，团队在真机逐项验证", "v0.4 界面与设置"],
        ["第 11 周", "自测补齐、发布单文件、文档与演示材料", "AI 生成自测用例与文档初稿", "68.5 MB 单文件、文档、截图"],
    ],
    widths=[2.4, 5.4, 4.2, 3.5],
)
body(
    "过程中最有价值的一条经验是“先定验收标准，再让 AI 写代码”。"
    "例如“回退”这个功能，如果只说“实现一键回退”，AI 给出的实现往往只覆盖顺利路径；"
    "而把验收标准写清楚——“覆盖过的文件回退后内容必须与原文件一致”“清理备份”“中途失败要能补偿”——"
    "生成结果的质量差异非常明显。"
)

heading("附录 D　迭代优化", 1)

heading("D.1 版本演进", 2)
caption("表 D1　主要版本与变化")
add_table(
    ["版本", "变化", "触发原因"],
    [
        ["v0.1", "建立核心层与批量重命名：只读规划、事务执行、日志、回退", "先把“安全底座”做对，再谈功能数量"],
        ["v0.2", "接入模块 2—6：转换、去重、摘要、导出、分类；统一日志 Kind 字段", "用户最期待的功能集中在文件读取与处理、格式转换与批量处理"],
        ["v0.3", "界面重构：左侧导航图标改为矢量线稿，统一页面骨架与状态提示", "初版图标用汉字占位，视觉不统一且显得临时"],
        ["v0.4", "新增完整设置页与常用设置面板、动效等级、页面记忆、默认值", "用户反馈“每次都要重填参数”，动效也缺少可控开关"],
    ],
    widths=[2.0, 8.4, 5.1],
)

heading("D.2 典型问题与修复", 2)
caption("表 D2　开发中遇到的主要问题及处理")
add_table(
    ["问题现象", "原因", "处理方式"],
    [
        ["中文 Markdown 导出的标题层级在再次打开时丢失", "导出时只写文本、未写入对应的标题样式", "写入 Normal 与 Heading1～3 样式定义，并按层级套用"],
        ["大批量预览时界面卡顿", "表格一次性渲染全部行", "启用行列虚拟化，只渲染可见区域"],
        ["仅大小写改名失败", "Windows 文件系统对大小写不敏感的写入顺序冲突", "识别为“仅大小写变化”，用临时文件中转完成"],
        ["PDF 中文显示为方块或错字", "依赖外部字体解析，中文字形缺失", "改为本地 WPF 排版后按页渲染，再封装 PDF"],
        ["转换中断后残留半截输出", "直接写入最终路径", "先写临时文件，成功后改名；失败或取消整批清理"],
        ["窗口在 200% 缩放下被裁切", "按像素写死布局", "Per-Monitor V2 + 设计画布 + Viewbox 等比适配"],
        ["AI 生成的实现引入联网能力", "默认实现倾向使用在线接口", "在 Prompt 与验收清单中写明“不得引用 HttpClient / 不接任何外部 API”，并在评审中删除"],
    ],
    widths=[5.2, 4.4, 5.9],
)

heading("附录 E　测试验证", 1)

heading("E.1 自动化自测", 2)
body(
    "团队编写了独立的核心自测程序（tests/LocalCraft.SelfTest），在临时目录中构造真实文件并验证"
    "规划、执行、日志、回退与算法行为。自测不依赖界面，可在命令行一键运行，"
    "2026 年 10 月 20 日的测试结果为 17 组用例、126 项断言、0 项失败。"
)
caption("表 E1　自测用例与覆盖点")
add_table(
    ["用例", "覆盖点"],
    [
        ["预览只读", "预览阶段不创建、不修改、不重命名任何文件"],
        ["重命名与回退", "新旧文件名、内容不丢、日志状态 Completed → Undone、备份清理"],
        ["覆盖与回退", "预览不覆盖；执行时备份目标；回退还原被覆盖文件内容"],
        ["仅大小写改名", "标记为仅大小写变化，可通过临时文件中转完成并可回退"],
        ["冲突检测", "多文件指向同一目标全部标记冲突，并给出原因"],
        ["非法文件名", "空主体、非法字符、保留名等被拒绝"],
        ["日志持久化", "日志写入后可重新加载，字段与状态一致"],
        ["文档转换与回退", "Markdown → DOCX → Markdown 中文内容回读一致，回退只删输出"],
        ["图片转换与回退", "PNG → JPG 生成有效图像文件，回退删除输出"],
        ["转换规划", "输出路径、同名序号、引擎匹配与不支持格式的标记"],
        ["PDF 导出", "中文内容导出为有效 PDF，页数与页面尺寸符合 A4"],
        ["多格式导出", "同一源文件同时生成 DOCX 与 PDF"],
        ["去重与摘要", "相似文档被分到同一组；关键句与关键词可定位到原文偏移"],
        ["分类与回退", "按规则生成目标子文件夹并移动，回退恢复原路径"],
        ["设置持久化", "设置写入 JSON 后可读回，默认值生效"],
        ["查找替换", "区分大小写、包含扩展名等规则的替换结果正确"],
        ["备份同步", "备份目录与日志记录保持一致，无遗留文件"],
    ],
    widths=[4.4, 11.0],
)

heading("E.2 功能验收", 2)
body("对照立项时写下的验收标准，逐项在真机上验证：")
caption("表 E2　验收标准与结果")
add_table(
    ["验收标准", "结果", "验证方式"],
    [
        ["Windows 电脑上双击 exe 即可运行", "通过", "在未安装 .NET 的环境运行 publish/win-x64/LocalCraft.exe"],
        ["断网状态下核心功能正常可用", "通过", "断网后逐模块执行重命名、转换、去重、摘要、导出、分类"],
        ["批量重命名后可通过日志一键回退", "通过", "自测用例 + 真机操作 300 个文件后回退"],
        ["能识别内容相似但文件名不同的文档", "通过", "同一文档改名与改格式后仍被分到同一相似组"],
        ["中文 Markdown 导出后标题层级与样式正确", "通过", "导出后用 Word 打开检查标题 1～3 级与段落间距"],
        ["格式转换与分类结果纳入回退", "通过", "回退中心删除输出、移动文件回原位"],
        ["4 GB 内存可流畅运行", "通过", "在 4 GB 设备上处理 2 000 个文件，界面无长时间卡死"],
        ["设置持久化与动画开关", "通过", "修改设置后重启，动效开关与默认值均生效"],
    ],
    widths=[6.2, 1.8, 7.4],
)

heading("E.3 边界与异常测试", 2)
body(
    "除正常路径外，团队专门验证了容易出问题的地方："
    "文件被其它程序占用时执行应整体失败并回滚，而不是留下半截结果；"
    "目标目录只读时应给出明确原因；超长路径与中文路径应正常处理；"
    "重命名链（A→B、B→A）应被判为冲突而不是隐式解决；"
    "去重遇到无文本层的扫描版 PDF 应跳过并在界面说明；"
    "转换任务在中途取消后不应残留临时文件。上述场景均已按预期处理并纳入自测。"
)

heading("附录 F　人工核查情况", 1)
body(
    "团队对本作品的全部 AI 生成内容做了逐项人工核查，核查记录如下。"
)
caption("表 F1　人工核查记录")
add_table(
    ["核查项", "核查方法", "结论"],
    [
        ["产品需求与缺口判断", "回溯问卷原始数据，逐条核对报告中引用的数字", "数据与报告一致；结论由团队独立判断"],
        ["代码正确性", "全量构建 + 17 组自测 + 真实文件夹端到端操作", "0 编译警告、0 断言失败"],
        ["离线与隐私", "审查工程引用与运行期网络行为，断网运行全部模块", "无 HTTP 客户端、无遥测、无外部 API 调用"],
        ["数据安全", "检查预览阶段写入、覆盖前备份、回退补偿逻辑", "预览零写入；覆盖先备份；回退可补偿"],
        ["界面与可达性", "在 100%/150%/200% 缩放下实机检查", "无横向裁切，文本物理尺寸可读"],
        ["中文排版", "导出后用 Word 与 PDF 阅读器逐个查看标题、段距、列表、上下标", "标题层级与基础样式正确；PDF 文字不可选中已如实说明"],
        ["文档一致性", "交叉核对文档中的模块状态、体积、测试数量与工程实际", "与 v0.4 实际交付一致"],
        ["学术诚信", "确认未使用他人代码与素材，AI 生成内容已声明", "符合大赛诚信与 AI 使用规范要求"],
    ],
    widths=[4.0, 6.6, 4.9],
)
body(
    "核查中发现并修正的问题主要有三类：AI 生成的实现默认引入联网或第三方库，已全部移除；"
    "文档初稿中“已支持 RAW”的表述超出实际能力，已改为“按系统解码器尝试并明确失败原因”；"
    "早期版本把测试数量写为旧值，已按最新自测结果更新。"
)

heading("附录 G　附加材料清单", 1)
body("随项目文档一同提交的材料如下，便于评审复现与体验：")
dash_item("可运行程序：publish/win-x64/LocalCraft.exe（自包含单文件，68.5 MB，双击即可运行）。", bold_prefix="① ")
dash_item("源代码：src/ 下三个程序集与 tests/ 下自测工程，附 README.md 与 docs/ARCHITECTURE.md。", bold_prefix="② ")
dash_item("界面截图：docs/screenshots/ 下各模块实际运行截图（本文档图 1—图 5 引自该目录）。", bold_prefix="③ ")
dash_item("调研材料：本地工具软件用户需求调研原始数据（CSV）与分析报告（DOCX）。", bold_prefix="④ ")
dash_item("演示视频与汇报 PPT：按大赛要求另行提交（MP4 ≤100 MB，PPT ≤20 页）。", bold_prefix="⑤ ")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("LocalCraft 团队")
run_font(r, SONG, 12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("2026 年 11 月")
run_font(r, SONG, 12)

# ---------------------------------------------------------------- 页眉页脚

sections = doc.sections
for idx, s in enumerate(sections):
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    s.different_first_page_header_footer = False
    if idx == 0:
        # 封面：页眉页脚留空
        for para in list(s.header.paragraphs):
            para.text = ""
        for para in list(s.footer.paragraphs):
            para.text = ""
        continue
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("杭州电子科技大学第二届 AIGC 应用创新大赛 · 项目文档 · LocalCraft 本地文件工坊")
    run_font(hr, SONG, 9, color=RGBColor(0x59, 0x59, 0x59))
    ppr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    ppr.append(borders)

    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    if idx == len(sections) - 1:
        r1 = fp.add_run("— ")
        run_font(r1, SONG, 9)
        add_page_field(fp)
        r2 = fp.add_run(" —")
        run_font(r2, SONG, 9)

# 打开文档时自动更新域（目录页码）
settings = doc.settings.element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)

doc.save(OUT)
print("saved:", OUT)
print("sections:", len(doc.sections), "paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
