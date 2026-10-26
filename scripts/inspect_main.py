import sys
from docx import Document
from docx.shared import Pt

def dump(path):
    print("=" * 100)
    print("FILE:", path)
    d = Document(path)
    print("sections:", len(d.sections))
    s = d.sections[0]
    print("page:", s.page_width, s.page_height, "margins:", s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)
    print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables))
    for i, p in enumerate(d.paragraphs):
        txt = p.text
        style = p.style.name
        sizes = {r.font.size.pt if r.font.size else None for r in p.runs}
        names = {r.font.name for r in p.runs}
        bold = {r.bold for r in p.runs}
        print(f"[{i:03d}] style={style!r} align={p.alignment} sz={sizes} font={names} b={bold} | {txt}")
    for ti, t in enumerate(d.tables):
        print(f"--- table {ti}: {len(t.rows)}x{len(t.columns)}")
        for r in t.rows:
            print("   | " + " | ".join(c.text.replace("\n", " ") for c in r.cells))

for p in sys.argv[1:]:
    dump(p)
